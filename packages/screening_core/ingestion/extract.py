"""Text extraction for PDF, DOCX and TXT with page/hidden-text heuristics."""
from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO

PARSER_VERSION = "extract-2.0.0"


@dataclass
class Extraction:
    text: str
    page_count: int | None
    image_only: bool
    hidden_text_suspected: bool
    tables_or_columns_suspected: bool
    notes: list[str] = field(default_factory=list)


def extract(extension: str, data: bytes) -> Extraction:
    if extension == ".txt":
        text = data.decode("utf-8")
        return Extraction(text, None, False, False, False)
    if extension == ".docx":
        return _docx(data)
    if extension == ".pdf":
        return _pdf(data)
    raise ValueError(f"unsupported extension {extension}")


def _docx(data: bytes) -> Extraction:
    import docx  # python-docx

    d = docx.Document(BytesIO(data))
    parts, hidden = [], False
    for p in d.paragraphs:
        for r in p.runs:
            if r.font.hidden or (r.font.color is not None and r.font.color.rgb is not None and str(r.font.color.rgb) == "FFFFFF") or (r.font.size is not None and r.font.size.pt < 3):
                hidden = True
        parts.append(p.text)
    tables = len(d.tables) > 0
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    text = "\n".join(parts)
    notes = ["tables present; reading order may be affected"] if tables else []
    return Extraction(text, None, len(text.strip()) == 0, hidden, tables, notes)


def _pdf(data: bytes) -> Extraction:
    import fitz  # PyMuPDF

    doc = fitz.open(stream=data, filetype="pdf")
    parts, hidden, columns, images = [], False, False, 0
    for page in doc:
        d = page.get_text("dict")
        images += len(page.get_images(full=True))
        xs = []
        for block in d.get("blocks", []):
            if block.get("type") != 0:
                continue
            xs.append(round(block["bbox"][0]))
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    if span.get("size", 10) < 3 or span.get("color", 0) == 0xFFFFFF:
                        hidden = True
        if len({x // 40 for x in xs}) >= 3 and len(xs) > 6:
            columns = True
        parts.append(page.get_text("text"))
    text = "\n".join(parts)
    image_only = len(text.strip()) < 40 and images > 0
    notes = []
    if columns:
        notes.append("multi-column layout suspected; reading order may be affected")
    return Extraction(text, doc.page_count, image_only, hidden, columns, notes)
